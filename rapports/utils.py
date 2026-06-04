from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from io import BytesIO
from datetime import date


def generer_pdf_ouvrier(ouvrier, salaire, paiements, mois, annee):
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                            rightMargin=2*cm, leftMargin=2*cm,
                            topMargin=2*cm, bottomMargin=2*cm)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('title',
        fontSize=18, alignment=TA_CENTER,
        fontName='Helvetica-Bold', spaceAfter=10)
    subtitle_style = ParagraphStyle('subtitle',
        fontSize=12, alignment=TA_CENTER,
        textColor=colors.grey, spaceAfter=20)
    section_style = ParagraphStyle('section',
        fontSize=13, fontName='Helvetica-Bold',
        textColor=colors.HexColor('#6C63FF'), spaceAfter=8, spaceBefore=15)

    elements = []

    # Titre
    elements.append(Paragraph("WorkManager", title_style))
    elements.append(Paragraph(
        f"Rapport Mensuel — {mois}/{annee}", subtitle_style))
    elements.append(Spacer(1, 0.5*cm))

    # Infos ouvrier
    elements.append(Paragraph("Informations de l'ouvrier", section_style))
    infos = [
        ['Nom complet', f"{ouvrier.nom} {ouvrier.prenom}"],
        ['Métier', ouvrier.get_metier_display()],
        ['Téléphone', ouvrier.telephone or '—'],
        ['Adresse', ouvrier.adresse or '—'],
        ['Prix journée', f"{ouvrier.prix_journee} DA"],
        ['Statut', ouvrier.get_statut_display()],
    ]
    t = Table(infos, colWidths=[5*cm, 12*cm])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#F3F3F3')),
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('PADDING', (0, 0), (-1, -1), 6),
        ('ROWBACKGROUNDS', (0, 0), (-1, -1),
         [colors.white, colors.HexColor('#FAFAFA')]),
    ]))
    elements.append(t)

    # Résumé pointage
    elements.append(Paragraph("Résumé du pointage", section_style))
    pointage_data = [
        ['Indicateur', 'Valeur'],
        ['Jours présents', str(salaire['jours_presents'])],
        ['Demi-journées', str(salaire['demi_journees'])],
        ['Absences', str(salaire['absents'])],
        ['Congés', str(salaire['conges'])],
        ['Total jours payés', str(salaire['total_jours_payes'])],
    ]
    t2 = Table(pointage_data, colWidths=[9*cm, 8*cm])
    t2.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#6C63FF')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('PADDING', (0, 0), (-1, -1), 6),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1),
         [colors.white, colors.HexColor('#FAFAFA')]),
    ]))
    elements.append(t2)

    # Résumé financier
    elements.append(Paragraph("Résumé financier", section_style))
    total_paye = sum(float(p.montant) for p in paiements)
    reste = salaire['montant_du'] - total_paye
    financier_data = [
        ['Élément', 'Montant'],
        ['Salaire dû', f"{salaire['montant_du']:.2f} DA"],
        ['Total payé', f"{total_paye:.2f} DA"],
        ['Reste à payer', f"{reste:.2f} DA"],
    ]
    t3 = Table(financier_data, colWidths=[9*cm, 8*cm])
    t3.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#6C63FF')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('PADDING', (0, 0), (-1, -1), 6),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#FFE4E4')
         if reste > 0 else colors.HexColor('#E4FFE4')),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
    ]))
    elements.append(t3)

    # Historique paiements
    if paiements:
        elements.append(Paragraph("Historique des paiements", section_style))
        paiement_data = [['Date', 'Montant', 'Note']]
        for p in paiements:
            paiement_data.append([
                str(p.date_paiement),
                f"{p.montant} DA",
                p.note or '—'
            ])
        t4 = Table(paiement_data, colWidths=[4*cm, 5*cm, 8*cm])
        t4.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#6C63FF')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('PADDING', (0, 0), (-1, -1), 6),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1),
             [colors.white, colors.HexColor('#FAFAFA')]),
        ]))
        elements.append(t4)

    # Footer
    elements.append(Spacer(1, 1*cm))
    elements.append(Paragraph(
        f"Généré le {date.today()} — WorkManager",
        ParagraphStyle('footer', fontSize=9,
                       textColor=colors.grey, alignment=TA_CENTER)
    ))

    doc.build(elements)
    buffer.seek(0)
    return buffer


def generer_pdf_global(ouvriers_data, mois, annee):
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                            rightMargin=1.5*cm, leftMargin=1.5*cm,
                            topMargin=2*cm, bottomMargin=2*cm)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('title',
        fontSize=18, alignment=TA_CENTER,
        fontName='Helvetica-Bold', spaceAfter=10)
    subtitle_style = ParagraphStyle('subtitle',
        fontSize=12, alignment=TA_CENTER,
        textColor=colors.grey, spaceAfter=20)

    elements = []
    elements.append(Paragraph("WorkManager", title_style))
    elements.append(Paragraph(
        f"Rapport Global — {mois}/{annee}", subtitle_style))
    elements.append(Spacer(1, 0.5*cm))

    # Tableau global
    headers = ['Nom', 'Métier', 'J.Présent',
               'Demi-J', 'J.Payés', 'Salaire Dû', 'Payé', 'Reste']
    data = [headers]

    total_du = 0
    total_paye = 0

    for o in ouvriers_data:
        data.append([
            f"{o['nom']} {o['prenom']}",
            o['metier'],
            str(o['jours_presents']),
            str(o['demi_journees']),
            str(o['total_jours_payes']),
            f"{o['montant_du']:.0f}",
            f"{o['total_paye']:.0f}",
            f"{o['reste_a_payer']:.0f}",
        ])
        total_du += o['montant_du']
        total_paye += o['total_paye']

    # Ligne total
    data.append([
        'TOTAL', '', '', '', '',
        f"{total_du:.0f} DA",
        f"{total_paye:.0f} DA",
        f"{total_du - total_paye:.0f} DA",
    ])

    col_widths = [4*cm, 3*cm, 2*cm, 2*cm, 2*cm, 2.5*cm, 2.5*cm, 2.5*cm]
    t = Table(data, colWidths=col_widths)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#6C63FF')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('PADDING', (0, 0), (-1, -1), 5),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2),
         [colors.white, colors.HexColor('#FAFAFA')]),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#E8E8FF')),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
    ]))
    elements.append(t)

    elements.append(Spacer(1, 1*cm))
    elements.append(Paragraph(
        f"Généré le {date.today()} — WorkManager",
        ParagraphStyle('footer', fontSize=9,
                       textColor=colors.grey, alignment=TA_CENTER)
    ))

    doc.build(elements)
    buffer.seek(0)
    return buffer


def generer_excel_global(ouvriers_data, mois, annee):
    buffer = BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Rapport {mois}-{annee}"

    # Style header
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="6C63FF")
    header_align = Alignment(horizontal="center")

    headers = ['Nom', 'Prénom', 'Métier', 'Prix/Jour',
               'J.Présents', 'Demi-J', 'Absences',
               'J.Payés', 'Salaire Dû', 'Total Payé', 'Reste']

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    # Données
    for row, o in enumerate(ouvriers_data, 2):
        ws.cell(row=row, column=1, value=o['nom'])
        ws.cell(row=row, column=2, value=o['prenom'])
        ws.cell(row=row, column=3, value=o['metier'])
        ws.cell(row=row, column=4, value=o['prix_journee'])
        ws.cell(row=row, column=5, value=o['jours_presents'])
        ws.cell(row=row, column=6, value=o['demi_journees'])
        ws.cell(row=row, column=7, value=o['absents'])
        ws.cell(row=row, column=8, value=o['total_jours_payes'])
        ws.cell(row=row, column=9, value=o['montant_du'])
        ws.cell(row=row, column=10, value=o['total_paye'])
        ws.cell(row=row, column=11, value=o['reste_a_payer'])

        # Colorier reste en rouge si > 0
        if o['reste_a_payer'] > 0:
            ws.cell(row=row, column=11).fill = PatternFill(
                "solid", fgColor="FFE4E4")

    # Largeurs colonnes
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 14

    wb.save(buffer)
    buffer.seek(0)
    return buffer